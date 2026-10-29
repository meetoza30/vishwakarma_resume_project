import os
import hashlib
import re
import pdfplumber
import docx
from fastapi import HTTPException, status
from core.config import settings
from services.ocr_service import extract_pdf_with_ocr_fallback


def compute_file_hash(content_bytes: bytes) -> str:
    """Compute SHA-256 hash of file contents for duplicate detection."""
    return hashlib.sha256(content_bytes).hexdigest()


def fix_mojibake(text: str) -> str:
    """Fix double-encoded UTF-8 characters."""
    if not text:
        return ""
    replacements = [
        ('\u00e2\u20ac\u201c', '\u2013'),
        ('\u00e2\u20ac\u201d', '\u2014'),
        ('\u00e2\u20ac\u2122', '\u2019'),
        ('\u00e2\u20ac\u0153', '\u201c'),
        ('\u00e2\u20ac\xa2', '\u2022'),
    ]
    for bad, good in replacements:
        text = text.replace(bad, good)
    return text


def validate_file_upload(filename: str, content_bytes: bytes) -> None:
    """
    Validate uploaded file:
    - Extension check (.pdf, .docx, .doc)
    - File size check (max 15MB)
    - Empty file check
    - Password protection / encryption check
    """
    ext = os.path.splitext(filename.lower())[1]
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Allowed formats: {', '.join(settings.ALLOWED_EXTENSIONS)}"
        )

    if len(content_bytes) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File '{filename}' is empty (0 bytes)."
        )

    if len(content_bytes) > settings.MAX_FILE_SIZE_BYTES:
        max_mb = settings.MAX_FILE_SIZE_BYTES / (1024 * 1024)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File size exceeds maximum limit of {max_mb:.0f} MB."
        )

    # Check PDF password protection
    if ext == ".pdf":
        if b"/Encrypt" in content_bytes[:4096] or b"/Encrypted" in content_bytes[:4096]:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"File '{filename}' is password-protected or encrypted. Please remove password protection before uploading."
            )


def extract_text_robust(file_path: str) -> str:
    """
    Robust text extraction pipeline supporting PDF, DOCX, and scanned image PDFs via OCR.
    """
    filename = os.path.basename(file_path).lower()
    ext = os.path.splitext(filename)[1]

    if ext in ['.docx', '.doc']:
        try:
            doc = docx.Document(file_path)
            full_text = []
            for p in doc.paragraphs:
                if p.text.strip():
                    full_text.append(p.text.strip())
            for t in doc.tables:
                for row in t.rows:
                    row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        full_text.append(row_txt)
            return fix_mojibake("\n".join(full_text))
        except Exception as e:
            print(f"Error extracting DOCX: {e}")
            raise HTTPException(status_code=400, detail=f"Corrupted or invalid DOCX file: {e}")

    elif ext == '.pdf':
        try:
            raw_text = extract_pdf_with_ocr_fallback(file_path)
            return fix_mojibake(raw_text)
        except Exception as e:
            print(f"Error extracting PDF: {e}")
            raise HTTPException(status_code=400, detail=f"Corrupted or unreadable PDF file: {e}")

    return ""
