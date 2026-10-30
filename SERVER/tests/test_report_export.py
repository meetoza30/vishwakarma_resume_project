import csv
import sys
from pathlib import Path

import openpyxl
import pdfplumber


SERVER_ROOT = Path(__file__).resolve().parents[1]
if str(SERVER_ROOT) not in sys.path:
    sys.path.insert(0, str(SERVER_ROOT))

from services.report_export import (  # noqa: E402
    EXCEL_REPORT_HEADERS,
    TABULAR_REPORT_HEADERS,
    build_tabular_report_row,
    dated_report_filename,
    aicte_eligible_role,
    write_excel_report,
    write_tabular_report,
)
from utils.generatereport import generate_match_report_pdf  # noqa: E402


def _result():
    return {
        "success": True,
        "target_position": "Assistant Professor",
        "candidate_name": "Anuradha Rai",
        "contact_number": "+91-9118062299",
        "email_id": "anuradha2002rai@gmail.com",
        "overall_score": 44,
        "breakdown": {
            "research_publications": {"score": 65},
            "experience_position_norms": {"score": 72},
            "patents_grants": {"score": 30},
            "leadership_roles": {"score": 40},
            "professional_dev": {"score": 55},
            "education": {"score": 80},
        },
        "experience_data": {"details": "Teaching/Research: 3.3 Yrs, Industry: 0 Yrs (Total: 3.3 Yrs)"},
        "qualification": {"formatted_string": "UG: B.Tech | PG: M.Tech | Ph.D.: None"},
        "research_areas": ["Deep Learning", "Intrusion Detection"],
        "previous_employers": ["Institute A", "Institute B"],
        "total_publications": 3,
        "total_indexed_publications": 2,
        "patents": "Nil",
        "grants_funding": "Nil",
    }


def test_report_row_matches_requested_schema_and_order():
    row = build_tabular_report_row(_result())
    assert len(row) == len(TABULAR_REPORT_HEADERS)
    assert row == [
        "Assistant Professor",
        "Anuradha Rai",
        "+91-9118062299",
        "anuradha2002rai@gmail.com",
        44,
        65,
        72,
        30,
        40,
        55,
        80,
        "Teaching/Research: 3.3 Yrs, Industry: 0 Yrs (Total: 3.3 Yrs)",
        "UG: B.Tech | PG: M.Tech | Ph.D.: None",
        "Deep Learning, Intrusion Detection",
        3,
        2,
        "Nil",
        "Nil",
    ]


def test_csv_header_and_row_have_identical_column_count(tmp_path):
    output = tmp_path / "faculty_report.csv"
    write_tabular_report(output, [_result()])

    with output.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))

    assert rows[0] == TABULAR_REPORT_HEADERS
    assert len(rows[0]) == 18
    assert len(rows[1]) == len(rows[0])
    assert rows[1][0] == "Assistant Professor"
    assert rows[1][15] == "2"


def test_csv_candidates_are_sorted_by_overall_score_descending(tmp_path):
    low = _result()
    low["candidate_name"] = "Lower Score"
    low["overall_score"] = 20
    high = _result()
    high["candidate_name"] = "Higher Score"
    high["overall_score"] = 95
    output = tmp_path / "sorted_report.csv"

    write_tabular_report(output, [low, high])
    with output.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert [rows[1][1], rows[2][1]] == ["Higher Score", "Lower Score"]


def test_failed_result_still_has_stable_column_count():
    row = build_tabular_report_row({"success": False, "error": "Failed processing"})
    assert len(row) == len(TABULAR_REPORT_HEADERS)
    assert row[1] == "Error: Failed processing"


def test_excel_report_matches_template_columns_and_sorts_by_score(tmp_path):
    lower = _result()
    lower["candidate_name"] = "Lower Score Candidate"
    lower["overall_score"] = 44
    lower["scopus_post_phd_publications"] = 2
    lower["aicte_eligibility"] = {
        "target_position": "Assistant Professor",
        "primary_recommended_status": "Eligible - Assistant Professor",
        "assistant_professor": {"eligible": True},
        "target_result": {"eligible": True, "checks": []},
    }
    higher = _result()
    higher["candidate_name"] = "Higher Score Candidate"
    higher["overall_score"] = 91

    output = tmp_path / "faculty_report.xlsx"
    write_excel_report(output, [lower, higher])
    workbook = openpyxl.load_workbook(output, data_only=True)
    sheet = workbook.active

    assert sheet.max_column == 22  # A spacer column plus 21 report columns.
    assert [sheet.cell(9, col).value for col in range(2, 23)] == EXCEL_REPORT_HEADERS
    assert sheet["B8"].value is None
    assert all("AICTE Eligibility Role" not in str(sheet.cell(row, col).value)
               for row in range(1, sheet.max_row + 1) for col in range(1, sheet.max_column + 1))
    assert sheet.cell(10, 4).value == "Higher Score Candidate"
    assert sheet.cell(10, 8).value == 91
    assert sheet.cell(11, 4).value == "Lower Score Candidate"
    assert sheet.cell(11, 5).value == "Assistant Professor"
    assert sheet.cell(11, 6).value == "+91-9118062299"
    assert sheet.cell(11, 7).value == "anuradha2002rai@gmail.com"
    assert sheet.cell(11, 8).value == 44
    assert "Experience & Norms: 72/100" in sheet.cell(11, 9).value
    assert "Technical Skills" not in sheet.cell(11, 9).value
    assert "UG: B.Tech" in sheet.cell(11, 10).value
    assert "PG: M.Tech" in sheet.cell(11, 10).value
    assert "Ph.D.: None" in sheet.cell(11, 10).value
    assert sheet.cell(11, 11).value == "Deep Learning, Intrusion Detection"
    assert sheet.cell(11, 12).value == 2
    assert sheet.cell(11, 13).value == "Nil"
    assert sheet.cell(11, 14).value == "Nil"
    assert sheet.cell(11, 15).value == "Teaching/Research: 3.3 Yrs, Industry: 0 Yrs (Total: 3.3 Yrs)"
    assert sheet.cell(11, 16).value == "• Institute A\n• Institute B"
    assert sheet.cell(11, 17).value == "Assistant Professor"
    assert sheet.cell(11, 18).value is None  # Salary fields are official-use only.
    assert sheet.cell(11, 3).value is None  # Official scheduling is left blank.
    assert sheet.cell(11, 22).value is None  # Official remarks are left blank.
    assert sheet.freeze_panes is None
    assert sheet.page_setup.orientation == "landscape"


def test_excel_distinguishes_applied_role_from_failed_norm_eligibility(tmp_path):
    candidate = _result()
    candidate["target_position"] = "Associate Professor"
    candidate["aicte_eligibility"] = {
        "target_position": "Associate Professor",
        "primary_recommended_status": "Eligible - Assistant Professor",
        "assistant_professor": {"eligible": True},
        "associate_professor": {"eligible": False},
        "target_result": {
            "eligible": False,
            "status": "Not Eligible",
            "checks": [
                {"name": "Ph.D. completed", "status": "fail"},
                {"name": "Post-Ph.D. experience", "status": "fail"},
            ],
        },
    }
    output = tmp_path / "ineligible_associate.xlsx"

    write_excel_report(output, [candidate])
    sheet = openpyxl.load_workbook(output, data_only=True).active

    # E is Applied for Post; Q is the independently recommended eligible role.
    assert sheet.cell(10, 5).value == "Associate Professor"
    assert sheet.cell(10, 17).value == "Assistant Professor"
    assert aicte_eligible_role(candidate) == "Assistant Professor"


def test_aicte_role_strips_eligible_prefix_and_handles_no_eligible_role():
    assert aicte_eligible_role({
        "aicte_eligibility": {
            "primary_recommended_status": "Eligible - Professor",
        }
    }) == "Professor"
    assert aicte_eligible_role({
        "aicte_eligibility": {"primary_recommended_status": "Not Eligible - Needs Verification"}
    }) == "Not Eligible"
    assert aicte_eligible_role({}) == "Not Evaluated"


def test_pdf_report_shows_aicte_role_remarks_and_score_order(tmp_path):
    def wrapped(name, score, role):
        result = _result()
        result["candidate_name"] = name
        result["overall_score"] = score
        result["aicte_eligibility"] = {
            "primary_recommended_status": role,
            "assistant_professor": {"eligible": True},
            "associate_professor": {"eligible": "Associate Professor" in role},
            "target_result": {"eligible": True, "checks": []},
        }
        return {"filename": f"{name}.pdf", "success": True, "match": result}

    output = tmp_path / "match_report.pdf"
    generate_match_report_pdf(
        results=[
            wrapped("Lower Score Candidate", 42, "Eligible - Assistant Professor"),
            wrapped("Higher Score Candidate", 89, "Eligible - Associate Professor"),
        ],
        jd_data={"position": "Associate Professor", "job_title": "Associate Professor"},
        output_path=str(output),
    )

    with pdfplumber.open(output) as pdf:
        text = " ".join(page.extract_text() or "" for page in pdf.pages)
    normalized = " ".join(text.split())
    assert normalized.index("Higher Score Candidate") < normalized.index("Lower Score Candidate")
    assert "Match Report" in normalized
    assert "Role: Associate Professor" in normalized
    assert "Technical Skills" not in normalized
    assert "Eligible - Associate Professor" not in normalized
    assert "AICTE Eligibility: Associate Professor" in normalized
    assert "Comments / Remarks (For Official Use)" in normalized


def test_report_filenames_are_date_first_and_job_unique():
    from datetime import datetime

    stamp = datetime(2026, 10, 30, 12, 0)
    assert dated_report_filename("pdf", "pdf", now=stamp) == "30.10.2026_match_report_pdf.pdf"
    assert dated_report_filename("excel", "xlsx", "abcd1234", stamp) == "30.10.2026_match_report_excel_abcd1234.xlsx"
